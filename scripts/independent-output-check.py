"""Independently decode/parse synthetic browser downloads; no app success flags."""
from pathlib import Path
from io import BytesIO
import json, zipfile
from PIL import Image
from pypdf import PdfReader
import pypdfium2 as pdfium
from docx import Document
from openpyxl import load_workbook
root=Path('.artifacts');results=[]
for directory in sorted(root.glob('astra-*')):
  folder=directory.name
  for file in (root/folder/'downloads').glob('*'):
    try:
      data=file.read_bytes(); detail={}
      if data.startswith(b'%PDF'):
        pdf=PdfReader(BytesIO(data))
        if pdf.is_encrypted:
          wrong=bool(pdf.decrypt('wrong-password'));password=next((p for p in ['SoraQA2026!','SyntheticQA7631'] if pdf.decrypt(p)),None);correct=bool(password)
          if wrong or not correct:raise ValueError('Password invariant failed')
          detail.update(encrypted=True,wrongPasswordRejected=True)
        detail.update(type='PDF',pages=len(pdf.pages),dimensions=[[float(p.mediabox.width),float(p.mediabox.height)] for p in pdf.pages],textCharacters=sum(len(p.extract_text() or '') for p in pdf.pages))
        rendering=pdfium.PdfDocument(data,password=password if pdf.is_encrypted else None)
        for index in {0,len(rendering)-1}:
          page=rendering[index];bitmap=page.render(scale=.5);im=bitmap.to_pil();im.load();detail.setdefault('renderedPages',[]).append(index+1)
          if index==0:im.save(root/folder/f'{file.stem}-independent.png')
          bitmap.close();page.close()
        rendering.close()
      elif data.startswith(b'PK'):
        z=zipfile.ZipFile(BytesIO(data)); assert z.testzip() is None
        if 'word/document.xml' in z.namelist():
          doc=Document(BytesIO(data));detail.update(type='DOCX',paragraphs=len(doc.paragraphs),tables=len(doc.tables),images=len(doc.inline_shapes))
        elif 'xl/workbook.xml' in z.namelist():
          wb=load_workbook(BytesIO(data),read_only=True);detail.update(type='XLSX',sheets=wb.sheetnames,rows=[s.max_row for s in wb]);wb.close()
        else:
          detail.update(type='ZIP',entries=len(z.namelist()))
          for name in z.namelist():
            blob=z.read(name)
            if blob.startswith(b'%PDF'):assert len(PdfReader(BytesIO(blob)).pages)>0
            elif name.lower().endswith(('.jpg','.png','.webp')):im=Image.open(BytesIO(blob));im.load()
      elif file.suffix.lower() in ['.jpg','.jpeg','.png','.webp']:
        im=Image.open(BytesIO(data));im.load();detail.update(type=im.format,dimensions=list(im.size),mode=im.mode,exifEntries=len(im.getexif()))
        if 'A' in im.getbands():
          detail['alphaRange']=list(im.getchannel('A').getextrema())
          if detail['alphaRange'][1]==0:raise ValueError('Entire image is transparent')
      elif file.suffix.lower()=='.txt':detail.update(type='TXT',characters=len(data.decode('utf-8')))
      else:raise ValueError('Unknown file signature')
      results.append({'file':str(file),'status':'PASS','details':detail})
    except Exception as e:results.append({'file':str(file),'status':'FAIL','error':str(e)})
(root/'astra-independent.json').write_text(json.dumps(results,indent=2))
print(json.dumps({'checked':len(results),'passed':sum(r['status']=='PASS' for r in results),'failed':sum(r['status']=='FAIL' for r in results)}))
